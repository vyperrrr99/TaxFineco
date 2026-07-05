import pdfplumber
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Legge tutte le pagine
all_text = []
with pdfplumber.open(r'C:\Users\vperr\Downloads\25VE005509_301225_1528.pdf') as pdf:
    for page in pdf.pages:
        text = page.extract_text()
        if text:
            all_text.append(text)

lines = '\n'.join(all_text).split('\n')

# Dizionario marche
brands = {
    'POL': 'Police', 'JCV': 'Just Cavalli', 'TMB': 'Timberland',
    'HUB': 'Hugo Boss', 'PLD': 'Polaroid', 'KSP': 'Kate Spade',
    'PCA': 'Pierre Cardin', 'CKC': 'Calvin Klein Collection',
    'CKJ': 'Calvin Klein Jeans', 'CK': 'Calvin Klein', 'ETR': 'Etro',
    'JC': 'Jimmy Choo', 'MSS': 'Missoni', 'ESC': 'Escada',
    'NR': 'Nina Ricci', 'TRU': 'Trussardi', 'TBS': 'Ted Baker',
    'FLA': 'Fila', 'HCK': 'Hackett', 'FUR': 'Furla', 'LJ': 'Liu Jo',
    'BVL': 'Bvlgari', 'VAL': 'Valentino', 'CPD': 'Chopard',
    'GUS': 'Guess', 'MMX': 'Max Mara', 'ADS': 'Adidas',
    'CH': 'Carolina Herrera', 'SWR': 'Swarovski', 'JVS': 'John Varvatos',
    'MOC': 'Moschino', 'VER': 'Versace', 'PRA': 'Prada',
    'DG': 'Dolce & Gabbana', 'EA': 'Emporio Armani', 'GA': 'Giorgio Armani',
    'OAK': 'Oakley', 'RB': 'Ray-Ban', 'PER': 'Persol',
    'VO': 'Vogue', 'TF': 'Tom Ford',
}

product_re = re.compile(
    r'^(\S+)\s+'
    r'(.+?)\s+'
    r'(PZ|KG|MT|CF|LT)\s+'
    r'(\d+(?:[.,]\d+)?)\s+'
    r'(\d+(?:[.,]\d+)?)\s+'
    r'(\d+(?:[.,]\d+)?)\s+'
    r'(\d+)$'
)
material_re = re.compile(r'^([A-Z][A-Z0-9 /]+?)\s+(CN|IT|DE|FR|ES|US|JP|TW|HK|UK|PT)\s*$')

desc_detail_re = re.compile(
    r'^([A-Z]{2,4})\s+'
    r'(SUN|OPT|OPT/SUN)?\s*'
    r'([A-Z0-9/._-]+)\s+'
    r'([A-Z0-9/._-]+)\s+'
    r'(\d{2})\s+'
    r'(\d{1,2})\s+'
    r'(\d{3})'
    r'(?:\s+(.+))?$'
)

products = []
i = 0
while i < len(lines):
    line = lines[i].strip()
    m = product_re.match(line)
    if m:
        codice = m.group(1)
        descrizione_raw = m.group(2).strip()
        um = m.group(3)
        qty   = float(m.group(4).replace(',', '.'))
        price = float(m.group(5).replace(',', '.'))
        imp   = float(m.group(6).replace(',', '.'))
        iva   = int(m.group(7))

        materiale, coo = '', ''
        if i+1 < len(lines):
            nm = material_re.match(lines[i+1].strip())
            if nm:
                materiale = nm.group(1).strip()
                coo = nm.group(2).strip()
                i += 1

        cod_marca = tipo = modello = colore = calibro = ponte = asta = marca = ''
        dd = desc_detail_re.match(descrizione_raw)
        if dd:
            cod_marca = dd.group(1)
            tipo      = dd.group(2) or ''
            modello   = dd.group(3)
            colore    = dd.group(4)
            calibro   = dd.group(5)
            ponte     = dd.group(6)
            asta      = dd.group(7)
            residuo   = dd.group(8) or ''
            marca     = residuo.strip() if residuo else brands.get(cod_marca, cod_marca)
        else:
            parts = descrizione_raw.split()
            cod_marca = parts[0] if parts else ''
            marca = brands.get(cod_marca, '')

        products.append({
            'Codice EAN': codice,
            'Descrizione completa': descrizione_raw,
            'Cod. Marca': cod_marca,
            'Marca': marca,
            'Tipo': tipo,
            'Modello': modello,
            'Colore': colore,
            'Calibro': calibro,
            'Ponte': ponte,
            'Asta': asta,
            'Materiale': materiale,
            'Paese Origine': coo,
            'U.M.': um,
            'Quantita': qty,
            'Prezzo Unitario': price,
            'Importo': imp,
            'Perc IVA': iva,
        })
    i += 1

print(f'Totale righe prodotto: {len(products)}')

# Crea Excel
wb = Workbook()
ws = wb.active
ws.title = "Prodotti Fattura"

header_font  = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
header_fill  = PatternFill('solid', fgColor='1F4E79')
alt_fill     = PatternFill('solid', fgColor='DCE6F1')
total_fill   = PatternFill('solid', fgColor='E2EFDA')
total_font   = Font(name='Calibri', bold=True, size=11)
num_fmt      = '#,##0.00'
int_fmt      = '#,##0'
thin = Side(style='thin', color='AAAAAA')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

columns = [
    ('Codice EAN',           18),
    ('Descrizione completa', 45),
    ('Cod. Marca',            9),
    ('Marca',                22),
    ('Tipo',                  8),
    ('Modello',              18),
    ('Colore',               10),
    ('Calibro',               8),
    ('Ponte',                 7),
    ('Asta',                  7),
    ('Materiale',            14),
    ('Paese Origine',        12),
    ('U.M.',                  6),
    ('Quantita',              9),
    ('Prezzo Unitario',      14),
    ('Importo',              13),
    ('Perc IVA',              8),
]

# Info fattura
ws['A1'] = 'FATTURA Nr. 25VE005509'
ws['A1'].font = Font(name='Calibri', bold=True, size=13, color='1F4E79')
ws['A2'] = 'Data: 30/12/2025  |  Fornitore: AERIAL VISION INTERNATIONAL S.P.A.  |  Cliente: YOCABE SRL'
ws['A2'].font = Font(name='Calibri', size=10, color='595959')
ws.merge_cells('A1:Q1')
ws.merge_cells('A2:Q2')
ws['A1'].alignment = Alignment(horizontal='left', vertical='center')
ws['A2'].alignment = Alignment(horizontal='left', vertical='center')
ws.row_dimensions[1].height = 22
ws.row_dimensions[2].height = 16

# Header
for col_idx, (col_name, col_width) in enumerate(columns, start=1):
    cell = ws.cell(row=4, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = border
    ws.column_dimensions[get_column_letter(col_idx)].width = col_width
ws.row_dimensions[4].height = 30

col_names = [c[0] for c in columns]
for row_idx, prod in enumerate(products, start=5):
    is_alt = (row_idx % 2 == 0)
    for col_idx, col_name in enumerate(col_names, start=1):
        val = prod.get(col_name, '')
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.font = Font(name='Calibri', size=10)
        cell.border = border
        cell.alignment = Alignment(vertical='center')
        if is_alt:
            cell.fill = alt_fill
        if col_name == 'Quantita':
            cell.number_format = int_fmt
            cell.alignment = Alignment(horizontal='right', vertical='center')
        elif col_name in ('Prezzo Unitario', 'Importo'):
            cell.number_format = num_fmt
            cell.alignment = Alignment(horizontal='right', vertical='center')
        elif col_name == 'Perc IVA':
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif col_name in ('Calibro', 'Ponte', 'Asta'):
            cell.alignment = Alignment(horizontal='center', vertical='center')

# Totali
tot_row = len(products) + 5
total_qty = sum(p['Quantita'] for p in products)
total_imp = sum(p['Importo'] for p in products)
iva_22    = sum(p['Importo'] * p['Perc IVA'] / 100 for p in products)

for c, v, fmt in [(1, 'TOTALE', None), (14, total_qty, int_fmt), (16, total_imp, num_fmt)]:
    cell = ws.cell(row=tot_row, column=c, value=v)
    cell.font = total_font
    cell.fill = total_fill
    cell.border = border
    if fmt:
        cell.number_format = fmt
        cell.alignment = Alignment(horizontal='right', vertical='center')
    else:
        cell.alignment = Alignment(horizontal='right', vertical='center')

r = tot_row + 1
cell = ws.cell(row=r, column=14, value='IVA 22%:')
cell.font = Font(name='Calibri', bold=True, size=10)
cell.alignment = Alignment(horizontal='right')
cell = ws.cell(row=r, column=16, value=round(iva_22, 2))
cell.number_format = num_fmt
cell.font = Font(name='Calibri', bold=True, size=10)

r = tot_row + 2
cell = ws.cell(row=r, column=14, value='TOTALE FATTURA:')
cell.font = Font(name='Calibri', bold=True, size=11, color='1F4E79')
cell.alignment = Alignment(horizontal='right')
cell = ws.cell(row=r, column=16, value=round(total_imp + iva_22, 2))
cell.number_format = num_fmt
cell.font = Font(name='Calibri', bold=True, size=11, color='1F4E79')

ws.freeze_panes = 'A5'
ws.auto_filter.ref = f'A4:{get_column_letter(len(columns))}{len(products)+4}'

out_path = r'C:\Users\vperr\Downloads\fattura_25VE005509.xlsx'
wb.save(out_path)
print(f'Salvato: {out_path}')
print(f'Righe: {len(products)}  |  Imponibile: {total_imp:.2f}  |  IVA: {iva_22:.2f}  |  Totale: {total_imp+iva_22:.2f}')
